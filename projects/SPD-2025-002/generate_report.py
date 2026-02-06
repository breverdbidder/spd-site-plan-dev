"""
Generate Bliss Palm Bay Feasibility Report (DOCX)
SPD-2025-002 — Sandy Pines Multifamily
"""

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

NAVY = RGBColor(0x1E, 0x3A, 0x5F)
RED = RGBColor(0xCC, 0x00, 0x00)
ORANGE = RGBColor(0xFF, 0x8C, 0x00)
YELLOW_DARK = RGBColor(0xB8, 0x86, 0x0B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)
LIGHT_GRAY = RGBColor(0xF2, 0xF2, 0xF2)

def set_cell_shading(cell, color_hex):
    """Set cell background color."""
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), color_hex)
    shading.set(qn('w:val'), 'clear')
    cell._tc.get_or_add_tcPr().append(shading)

def set_cell_border(cell, **kwargs):
    """Set cell borders."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge, val in kwargs.items():
        element = OxmlElement(f'w:{edge}')
        element.set(qn('w:val'), val.get('val', 'single'))
        element.set(qn('w:sz'), val.get('sz', '4'))
        element.set(qn('w:color'), val.get('color', '000000'))
        element.set(qn('w:space'), val.get('space', '0'))
        tcBorders.append(element)
    tcPr.append(tcBorders)

def add_styled_heading(doc, text, level=1):
    """Add a heading with navy color and Arial font."""
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = NAVY
        run.font.name = 'Arial'
    return heading

def add_body_text(doc, text, bold=False, italic=False):
    """Add body paragraph with Arial font."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(10)
    run.bold = bold
    run.italic = italic
    return p

def add_bullet(doc, text, bold_prefix=None):
    """Add a bullet point."""
    p = doc.add_paragraph(style='List Bullet')
    if bold_prefix:
        run_b = p.add_run(bold_prefix)
        run_b.font.name = 'Arial'
        run_b.font.size = Pt(10)
        run_b.bold = True
        run = p.add_run(text)
    else:
        run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(10)
    return p

def style_table_header(row, color_hex='1E3A5F'):
    """Style a table header row."""
    for cell in row.cells:
        set_cell_shading(cell, color_hex)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.color.rgb = WHITE
                run.font.name = 'Arial'
                run.font.size = Pt(9)
                run.bold = True

def style_table_cell(cell, text, align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, color=None, size=9):
    """Style a table cell."""
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = align
    run = p.add_run(str(text))
    run.font.name = 'Arial'
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color

def add_table(doc, headers, rows, col_widths=None):
    """Add a formatted table."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'

    # Header
    for i, h in enumerate(headers):
        style_table_cell(table.rows[0].cells[i], h, WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    style_table_header(table.rows[0])

    # Data rows
    for r_idx, row_data in enumerate(rows):
        for c_idx, val in enumerate(row_data):
            if isinstance(val, dict):
                style_table_cell(table.rows[r_idx + 1].cells[c_idx], val['text'],
                               val.get('align', WD_ALIGN_PARAGRAPH.LEFT),
                               val.get('bold', False),
                               val.get('color', None))
            else:
                style_table_cell(table.rows[r_idx + 1].cells[c_idx], str(val))
        if r_idx % 2 == 1:
            for cell in table.rows[r_idx + 1].cells:
                set_cell_shading(cell, 'F2F2F2')

    if col_widths:
        for i, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Inches(w)

    return table

def create_report():
    doc = Document()

    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(10)

    # Adjust margins
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # =========================================================================
    # COVER PAGE
    # =========================================================================
    for _ in range(6):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('BLISS PALM BAY')
    run.font.name = 'Arial'
    run.font.size = Pt(36)
    run.font.color.rgb = NAVY
    run.bold = True

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run('Site Plan Feasibility Report')
    run.font.name = 'Arial'
    run.font.size = Pt(20)
    run.font.color.rgb = NAVY

    doc.add_paragraph()

    details = [
        '2165 Sandy Pines Drive NE, Palm Bay, FL 32905',
        'Parcel ID: 28-37-27-00-7  |  BCPAO Account: 2835546',
        '',
        'Prepared for: The Property Squad, Inc.',
        'Prepared by: Everest Capital USA',
        'Date: February 6, 2026',
        '',
        'Project ID: SPD-2025-002',
        'Status: HOLD PENDING PUD REVIEW'
    ]
    for d in details:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(d)
        run.font.name = 'Arial'
        run.font.size = Pt(12)
        if 'HOLD' in d:
            run.font.color.rgb = RED
            run.bold = True

    doc.add_page_break()

    # =========================================================================
    # TABLE OF CONTENTS placeholder
    # =========================================================================
    add_styled_heading(doc, 'Table of Contents', level=1)
    toc_items = [
        'Executive Summary',
        'Section 1: Property Overview',
        'Section 2: Zoning & Land Use',
        'Section 3: Site Analysis',
        'Section 4: Environmental Constraints',
        'Section 5: Infrastructure & Utilities',
        'Section 6: Traffic & Access',
        'Section 7: HOA & CC&R Analysis',
        'Section 8: Development Scenarios',
        'Section 9: Cost Estimate',
        'Section 10: Permit Roadmap',
        'Section 11: Risk Matrix',
        'Section 12: Recommendation'
    ]
    for item in toc_items:
        p = doc.add_paragraph()
        run = p.add_run(item)
        run.font.name = 'Arial'
        run.font.size = Pt(11)

    doc.add_page_break()

    # =========================================================================
    # EXECUTIVE SUMMARY
    # =========================================================================
    add_styled_heading(doc, 'Executive Summary', level=1)

    p = doc.add_paragraph()
    run = p.add_run('RECOMMENDATION: HOLD / CONDITIONAL NO-GO')
    run.font.name = 'Arial'
    run.font.size = Pt(14)
    run.font.color.rgb = RED
    run.bold = True

    add_body_text(doc, '')
    add_body_text(doc,
        'This feasibility report evaluates a proposed multifamily development on a 1.114-acre vacant parcel '
        'at 2165 Sandy Pines Drive NE, Palm Bay, Florida. The parcel is zoned PUD (Planned Unit Development) '
        'with a High Density Residential (HDR) future land use designation, theoretically supporting up to '
        '20 dwelling units per acre (22 units maximum).')

    add_body_text(doc, '')
    add_body_text(doc, 'Key Metrics:', bold=True)

    metrics = [
        ('Land Value:', ' $55,000 (BCPAO 2025 market value)'),
        ('Acquisition Price:', ' $35,100 (tax deed, July 2023)'),
        ('Gross Area:', ' 1.114 acres (48,516 SF) per survey traverse calculation'),
        ('Net Buildable (3-story):', ' 11,325 SF (23.3%) after wellhead easement + setbacks'),
        ('Maximum Density:', ' 22 units gross / 11 units net (after easement reduction)'),
        ('Environmental Costs:', ' $77,000 - $203,000 (1.4x to 3.7x land value)'),
        ('Total Development Cost:', ' $3.79M (11 units) / $7.26M (22 units)'),
        ('Financial Feasibility:', ' NOT FEASIBLE — development cost exceeds stabilized value by ~2x'),
    ]
    for prefix, text in metrics:
        add_bullet(doc, text, bold_prefix=prefix)

    add_body_text(doc, '')
    add_body_text(doc, 'Five Critical Constraints Identified:', bold=True)

    constraints_exec = [
        ('1. HOA CC&Rs (POTENTIAL FATAL):', ' Sandy Pines Preserve is a gated community of exclusively 1-2 story single-family homes. CC&Rs almost certainly prohibit 3-4 story multifamily construction.'),
        ('2. No Sewer/Water Available (HIGH):', ' Palm Bay GIS confirms no sewer or water service at parcel. Sewer extension required at developer cost.'),
        ('3. FEMA Zone A Flood (HIGH):', ' Entire parcel in Special Flood Hazard Area with no BFE determined. Flood study required ($5K-$25K).'),
        ('4. NWI Wetland PEM1Cd (HIGH):', ' 10.3-acre freshwater emergent wetland envelops parcel. Mitigation credits: $40K-$100K.'),
        ('5. Wellhead Protection (HIGH):', ' 200-ft radius easement covers 47% of parcel. No construction in easement zone for ~10 years.'),
    ]
    for prefix, text in constraints_exec:
        add_bullet(doc, text, bold_prefix=prefix)

    add_body_text(doc, '')
    add_body_text(doc,
        'The combination of five overlapping constraints makes multifamily development economically infeasible '
        'at current market conditions. Environmental compliance costs alone ($77K-$203K) exceed the land value '
        '($55K). Total development cost for an 11-unit project ($3.79M) is approximately double the stabilized '
        'property value ($1.93M at a 6% cap rate). The project should be placed on HOLD pending CC&R retrieval '
        'and PUD ordinance review. If CC&Rs prohibit multifamily, the project should be KILLED for the intended use.')

    doc.add_page_break()

    # =========================================================================
    # SECTION 1: PROPERTY OVERVIEW
    # =========================================================================
    add_styled_heading(doc, 'Section 1: Property Overview', level=1)

    add_body_text(doc, 'Location & Legal Description', bold=True)
    prop_data = [
        ['Address', '2165 Sandy Pines Drive NE, Palm Bay, FL 32905'],
        ['Parcel ID', '28-37-27-00-7'],
        ['BCPAO Account', '2835546'],
        ['Legal Description', 'E 1/2 of S 1/2 of NE 1/4 of Section 27, T28S, R37E, less E 83 ft'],
        ['Plat References', 'PB 32/78, ORB 3636/2327, ORB 3731/2000'],
        ['Jurisdiction', 'City of Palm Bay, Brevard County, Florida'],
        ['Millage Code', '34U0 (Palm Bay)'],
    ]
    add_table(doc, ['Field', 'Value'], prop_data, [2.0, 5.0])

    add_body_text(doc, '')
    add_body_text(doc, 'Ownership', bold=True)
    owner_data = [
        ['Current Owner', 'The Property Squad, Inc.'],
        ['Owner Type', 'Florida Corporation'],
        ['Mailing Address', '390 Roosevelt Ave, Satellite Beach, FL 32937'],
        ['Acquisition Date', 'July 24, 2023'],
        ['Acquisition Price', '$35,100'],
        ['Deed Type', 'Tax Deed (XD) — Book 9844, Page 0688'],
        ['2025 Market Value', '$55,000'],
        ['2025 Taxable Value', '$55,000'],
        ['Property Use Code', '0008 — Vacant Residential (Multi-Family, Unplatted)'],
        ['Total Buildings', '0 (vacant land)'],
    ]
    add_table(doc, ['Field', 'Value'], owner_data, [2.0, 5.0])

    add_body_text(doc, '')
    add_body_text(doc, 'Sales History', bold=True)
    sales_data = [
        ['07/24/2023', '$35,100', 'Tax Deed (XD)', 'Book 9844/Page 0688'],
        ['06/30/2011', '$4,300', 'Tax Deed (XD)', 'Book 6419/Page 1735'],
        ['10/29/2004', '$277,000', 'Trustee Deed (TD)', 'Book 5382/Page 2335'],
    ]
    add_table(doc, ['Date', 'Price', 'Deed Type', 'Recording'], sales_data, [1.5, 1.2, 1.8, 2.5])

    add_body_text(doc, '')
    add_body_text(doc, 'Note: Two consecutive tax deed sales (2011 and 2023) indicate long-term vacancy. '
                  'The $277K-to-$35K price decline reflects the severe constraints on this parcel, not general '
                  'market conditions.', italic=True)

    doc.add_page_break()

    # =========================================================================
    # SECTION 2: ZONING & LAND USE
    # =========================================================================
    add_styled_heading(doc, 'Section 2: Zoning & Land Use', level=1)

    add_body_text(doc, 'Current Zoning Designation', bold=True)
    zoning_data = [
        ['Current Zoning', 'PUD (Planned Unit Development)'],
        ['Future Land Use', 'HDR (High Density Residential)'],
        ['Maximum Density (HDR)', 'Up to 20 du/acre'],
        ['PUD Type', 'Small PUD (SPUD) — min 1.0 acre (Section 185.061)'],
        ['PUD Ordinance', 'NOT YET OBTAINED — must request from City Clerk'],
        ['Zoning Source', 'Palm Bay GIS REST — live query Feb 6, 2026'],
    ]
    add_table(doc, ['Field', 'Value'], zoning_data, [2.5, 4.5])

    add_body_text(doc, '')
    add_body_text(doc, 'PUD Zoning Implications', bold=True)
    add_body_text(doc,
        'PUD zoning means development standards (setbacks, height, density, permitted uses) are controlled '
        'by the approved PUD ordinance and development plan — NOT by the standard RM district tables in the '
        'municipal code. The PUD ordinance for Sandy Pines Preserve has not been located through online searches. '
        'It must be obtained from the Palm Bay City Clerk. Without this document, no definitive feasibility '
        'determination can be made.')

    add_body_text(doc, '')
    add_body_text(doc, 'RM-15 Reference Standards (Fallback Benchmark)', bold=True)
    add_body_text(doc,
        'Since the PUD ordinance is unavailable, RM-15 standards are used as a reference benchmark. '
        'RM-15 is the most comparable standard zoning district for HDR future land use:')

    rm15_data = [
        ['Maximum Density', '15 du/acre'],
        ['Front Setback', '25 ft OR building height, whichever is greater'],
        ['Side Setback', '15 ft OR building height, whichever is greater'],
        ['Rear Setback', '25 ft OR building height, whichever is greater'],
        ['Max Building Height', '40 ft'],
        ['Building Separation', '30 ft'],
        ['Parking Setback', '10 ft from property line'],
        ['Open Space', '25% of gross site area (PUD requirement, Section 173.070)'],
    ]
    add_table(doc, ['Standard', 'Requirement'], rm15_data, [2.5, 4.5])

    add_body_text(doc, '')
    add_body_text(doc, 'Ordinance 2024-45 (Nov 7, 2024): All RM-10 districts rezoned to RM-15 citywide '
                  '(361 acres affected). This ordinance does NOT apply to PUD-zoned parcels.', italic=True)

    doc.add_page_break()

    # =========================================================================
    # SECTION 3: SITE ANALYSIS
    # =========================================================================
    add_styled_heading(doc, 'Section 3: Site Analysis', level=1)

    add_body_text(doc, 'Survey-Verified Dimensions', bold=True)
    add_body_text(doc, 'Source: Smith Survey Company (L.B. 7436), Scale 1" = 30\'')

    dim_data = [
        ['North Boundary', 'S89\u00b031\'14"E', '200.00 ft', 'Adjacent to Tract D (Preserve)'],
        ['East Boundary', 'S00\u00b028\'46"W', '277.15 ft', 'Adjacent to Tract O (Preserve)'],
        ['South (Sandy Pines Dr)', 'Curve R=1036.12\'', '214.67 ft (arc)', 'Road frontage (curved)'],
        ['West Boundary', 'N00\u00b028\'46"E', '200.00 ft', 'Adjacent to Tract D (Preserve)'],
    ]
    add_table(doc, ['Boundary', 'Bearing', 'Distance', 'Adjacent'], dim_data, [1.5, 1.8, 1.5, 2.2])

    add_body_text(doc, '')
    add_body_text(doc, 'Area Calculations', bold=True)
    area_data = [
        ['Computed Area (traverse)', '48,516 SF', '1.114 acres'],
        ['GIS Area', '48,424 SF', '1.11 acres'],
        ['BCPAO Recorded', '47,916 SF', '1.10 acres'],
        ['Perimeter', '891.82 ft', '(GIS = computed)'],
    ]
    add_table(doc, ['Source', 'Square Feet', 'Acres'], area_data, [2.5, 2.0, 2.5])

    add_body_text(doc, '')
    add_body_text(doc, 'Buildable Area Analysis', bold=True)

    build_data = [
        ['Total Parcel', '48,516', '1.114', '100%'],
        ['Wellhead Easement', '22,803', '0.524', '47%'],
        ['Net (excl. wellhead)', '25,713', '0.590', '53%'],
        ['Net after setbacks (2-story, 25ft)', '14,555', '0.334', '30.0%'],
        ['Net after setbacks (3-story, 35ft)', '11,325', '0.260', '23.3%'],
        ['Net after setbacks (4-story, 40ft)', '9,861', '0.226', '20.3%'],
    ]
    add_table(doc, ['Description', 'Square Feet', 'Acres', '% of Parcel'], build_data, [2.8, 1.3, 1.0, 1.2])

    add_body_text(doc, '')
    add_body_text(doc, 'Wellhead Protection Easement', bold=True)
    add_body_text(doc,
        'A 200-foot radius circular easement around a City of Palm Bay public water supply well covers the '
        'western 47% of the parcel (22,803 SF). Per Florida DEP Chapter 62-521 FAC, no new construction, '
        'impervious surfaces, or regulated substance storage is permitted within this zone. Surface parking '
        'may be allowed (survey annotates "PARKING" in this area). The well may be decommissioned in ~10 years '
        'per Palm Bay Utilities (Tim Roberts, Sept 11, 2025).')

    doc.add_page_break()

    # =========================================================================
    # SECTION 4: ENVIRONMENTAL CONSTRAINTS
    # =========================================================================
    add_styled_heading(doc, 'Section 4: Environmental Constraints', level=1)

    add_body_text(doc, 'FEMA Flood Zone', bold=True)
    flood_data = [
        ['Flood Zone', 'A (Special Flood Hazard Area)'],
        ['SFHA', 'Yes — 1% annual chance (100-year) flood'],
        ['Base Flood Elevation', 'NOT DETERMINED (Study Type: NP)'],
        ['DFIRM ID', '12009C'],
        ['CRS Rating', '7 (15% flood insurance discount)'],
        ['Flood Study Cost', '$5,000 - $25,000 (HEC-RAS hydraulic modeling)'],
    ]
    add_table(doc, ['Item', 'Detail'], flood_data, [2.5, 4.5])

    add_body_text(doc, '')
    add_body_text(doc,
        'The entire parcel is in FEMA Zone A with no published Base Flood Elevation. The developer must '
        'commission a hydraulic study (HEC-RAS modeling) to establish BFE before any building design can proceed. '
        'All habitable floors must be elevated above BFE plus 1 foot minimum freeboard per Florida Building Code.')

    add_body_text(doc, '')
    add_body_text(doc, 'National Wetlands Inventory (NWI)', bold=True)
    wetland_data = [
        ['NWI Code', 'PEM1Cd'],
        ['Type', 'Freshwater Emergent Wetland'],
        ['Feature Total', '10.316 acres (parcel within larger polygon)'],
        ['Classification', 'Palustrine, Emergent, Persistent, Seasonally Flooded, Partially Drained'],
        ['Delineation Cost', '$3,000 - $8,000'],
        ['Mitigation Credits', '0.3-0.5 UMAM credits at $125K-$200K/credit'],
        ['Mitigation Cost', '$40,000 - $100,000'],
        ['Primary Bank', 'Farmton Mitigation Bank (south Brevard service area)'],
    ]
    add_table(doc, ['Item', 'Detail'], wetland_data, [2.5, 4.5])

    add_body_text(doc, '')
    add_body_text(doc, 'Listed Species', bold=True)
    add_body_text(doc,
        'Palm Bay has an active USFWS-approved Habitat Conservation Plan (HCP) covering Florida Scrub-Jay, '
        'Eastern Indigo Snake, and Gopher Tortoise. The HCP scrub-jay fee is $125 per quarter-acre ($560 total). '
        'A 100% gopher tortoise survey is required before any land clearing. Wood stork foraging habitat '
        '(PEM1Cd wetland type) may trigger USFWS consultation.')

    add_body_text(doc, '')
    add_body_text(doc, 'Environmental Cost Summary', bold=True)
    env_cost_data = [
        ['Flood study / BFE determination', '$5,000', '$25,000'],
        ['Elevation certificates (3x)', '$500', '$6,000'],
        ['Wetland delineation', '$3,000', '$8,000'],
        ['Listed species surveys', '$2,000', '$8,000'],
        ['Environmental consultant (ERP)', '$15,000', '$40,000'],
        ['SJRWMD ERP application fee', '$9,000', '$9,000'],
        ['Wetland mitigation credits', '$40,000', '$100,000'],
        ['Gopher tortoise survey/relocation', '$1,500', '$6,000'],
        ['Scrub-jay HCP fee', '$560', '$560'],
        [{'text': 'TOTAL', 'bold': True}, {'text': '$76,560', 'bold': True}, {'text': '$202,560', 'bold': True}],
    ]
    add_table(doc, ['Item', 'Low Estimate', 'High Estimate'], env_cost_data, [3.5, 1.5, 1.5])

    add_body_text(doc, '')
    add_body_text(doc,
        'Environmental costs ($77K-$203K) represent 139%-368% of the land value ($55K). '
        'Wetland mitigation credits are the single largest cost component at $40K-$100K.', italic=True)

    doc.add_page_break()

    # =========================================================================
    # SECTION 5: INFRASTRUCTURE & UTILITIES
    # =========================================================================
    add_styled_heading(doc, 'Section 5: Infrastructure & Utilities', level=1)

    util_data = [
        ['Water', 'City of Palm Bay (PBUD)', 'NO*', 'Wellhead within 200ft but no main at property line'],
        ['Sewer', 'City of Palm Bay (PBUD)', 'NO', 'Not in SANC/WANC phase. Extension required.'],
        ['Electric', 'FPL', 'YES', 'Serves all Brevard County'],
        ['Natural Gas', 'Florida City Gas', 'LIKELY', 'Address-specific confirmation needed'],
        ['Stormwater', 'MTWCD / Palm Bay', 'N/A', 'On-site system required. MTWCD permit required.'],
        ['Cable/Internet', 'Spectrum', 'YES', '100% coverage, up to 1 Gbps'],
    ]
    add_table(doc, ['Utility', 'Provider', 'Available', 'Notes'], util_data, [1.2, 2.0, 0.8, 3.0])

    add_body_text(doc, '')
    add_body_text(doc, '* Water Availability Paradox: Despite a city wellhead within 200 feet, Palm Bay GIS '
                  'designates this parcel as Water_Avail: NO. The wellhead is a raw water production well '
                  'feeding a treatment plant, not the distribution system. A water main may not extend to '
                  'this parcel\'s frontage.', italic=True)

    add_body_text(doc, '')
    add_body_text(doc, 'Sewer Extension Analysis', bold=True)
    add_body_text(doc,
        'Adjacent parcels (2153, 2240, 2260, 2300 Sandy Pines Dr) have both water and sewer available, '
        'suggesting the sewer main runs along Sandy Pines Drive but does not extend to this parcel\'s frontage. '
        'Sewer extension cost is estimated at $50-$100 per linear foot. Distance to nearest main is unknown '
        'but may be at 2153 Sandy Pines Dr (adjacent parcel).')

    add_body_text(doc, '')
    add_body_text(doc, 'Impact Fees', bold=True)
    fee_data = [
        ['Palm Bay Water', '$2,221/unit'],
        ['Palm Bay Sewer', '$3,544/unit'],
        ['Brevard County Transportation', '$2,381/unit (3+ story)'],
        ['Brevard County Education', '$1,941/unit'],
        ['Brevard County Other', '$1,750/unit (est.)'],
        [{'text': 'Total per Unit', 'bold': True}, {'text': '$11,837/unit', 'bold': True}],
        ['Total (11 units)', '$130,207'],
        ['Total (22 units)', '$260,414'],
    ]
    add_table(doc, ['Fee Category', 'Amount'], fee_data, [3.5, 3.5])

    doc.add_page_break()

    # =========================================================================
    # SECTION 6: TRAFFIC & ACCESS
    # =========================================================================
    add_styled_heading(doc, 'Section 6: Traffic & Access', level=1)

    add_body_text(doc, 'Road Network', bold=True)
    add_body_text(doc,
        'Sandy Pines Drive NE is a PRIVATE road within the Sandy Pines Preserve gated community. '
        'All traffic from the proposed development must pass through the community\'s security gate '
        'and HOA-controlled roads before reaching Malabar Road (SR 514).')

    add_body_text(doc, '')
    add_body_text(doc, 'ITE Trip Generation (LU 221 — Mid-Rise Multifamily)', bold=True)
    trip_data = [
        ['11-unit scenario', '50 ADT', '4', '4'],
        ['22-unit scenario', '100 ADT', '8', '9'],
    ]
    add_table(doc, ['Scenario', 'Daily Trips', 'AM Peak', 'PM Peak'], trip_data, [2.0, 1.5, 1.5, 1.5])

    add_body_text(doc, '')
    add_body_text(doc,
        'Both scenarios generate well below 100 peak-hour trips, which is the common Florida threshold '
        'for requiring a formal Traffic Impact Study (TIS). The project\'s impact on Malabar Road (SR 514) '
        'is de minimis at 0.3-0.6% of corridor volume.')

    add_body_text(doc, '')
    add_body_text(doc, 'Malabar Road (SR 514) Conditions', bold=True)
    add_bullet(doc, ' 16,000 AADT on a 2-lane undivided road — severely congested', bold_prefix='Current:')
    add_bullet(doc, ' FDOT PD&E completed May 2018 — widen to 4-lane divided ($84M-$97M). NOT FUNDED.',
               bold_prefix='Widening:')
    add_bullet(doc, ' 56% increase in crashes over 5 years', bold_prefix='Safety:')

    add_body_text(doc, '')
    add_body_text(doc, 'Private Road / HOA Considerations', bold=True)
    add_bullet(doc, 'HOA Board approval required for any change in use or increased traffic')
    add_bullet(doc, 'Road maintenance contribution agreement likely needed')
    add_bullet(doc, 'Emergency vehicle access must be verified with Palm Bay Fire-Rescue')
    add_bullet(doc, 'Construction traffic access through gated community requires HOA coordination')

    doc.add_page_break()

    # =========================================================================
    # SECTION 7: HOA & CC&R ANALYSIS
    # =========================================================================
    add_styled_heading(doc, 'Section 7: HOA & CC&R Analysis', level=1)

    p = doc.add_paragraph()
    run = p.add_run('WARNING: POTENTIAL FATAL CONSTRAINT')
    run.font.name = 'Arial'
    run.font.size = Pt(12)
    run.font.color.rgb = RED
    run.bold = True

    add_body_text(doc, '')
    add_body_text(doc, 'Community Profile', bold=True)
    hoa_data = [
        ['Community', 'Sandy Pines Preserve'],
        ['Type', 'Gated community with security access'],
        ['Year Built', '1998-2003'],
        ['Housing Types', 'Single-family detached + townhomes (1-2 story only)'],
        ['Multifamily Buildings', 'NONE exist anywhere in community'],
        ['HOA Fee (Phase 1&2)', '$246/quarter (~$82/month)'],
        ['Median Sale Price', '$268,250 (Phase 1&2)'],
    ]
    add_table(doc, ['Item', 'Detail'], hoa_data, [2.5, 4.5])

    add_body_text(doc, '')
    add_body_text(doc, 'HOA Governance Structure', bold=True)
    hoa_gov = [
        ['Phases 1 & 2', 'Leland Management', '(407) 447-9955', 'N97000006305'],
        ['Phases 3 & 4', 'Artemis Lifestyles', 'N/A', 'N01000002407'],
        ['Courtyards', 'Sentry Management', '(407) 788-6700', 'N99000003017'],
    ]
    add_table(doc, ['Phase', 'Management', 'Phone', 'Sunbiz Filing'], hoa_gov, [1.5, 2.0, 1.5, 2.0])

    add_body_text(doc, '')
    add_body_text(doc, 'Multifamily Feasibility Under CC&Rs', bold=True)
    add_body_text(doc,
        'Based on extensive research, multifamily construction (3-4 story apartment buildings) is '
        'ALMOST CERTAINLY PROHIBITED by the Sandy Pines Preserve CC&Rs. Evidence includes:')

    prohib_items = [
        'All existing structures across all phases are exclusively 1-2 story single-family and townhome',
        'HOA incorporated under FL Chapter 720 (Homeowners Association Act)',
        'No apartments, condos, or multifamily buildings exist anywhere in the community',
        'Typical late-1990s gated community CC&Rs restrict to single-family/townhome use types',
        'Architectural Review Committee (ARC) approval almost certainly required',
    ]
    for item in prohib_items:
        add_bullet(doc, item)

    add_body_text(doc, '')
    add_body_text(doc, 'CC&R Status:', bold=True)
    add_body_text(doc,
        'The CC&Rs are recorded in OR Book 3636, Page 2327 at the Brevard County Clerk. They have NOT '
        'yet been obtained or reviewed. This is the SINGLE MOST CRITICAL action item. If CC&Rs prohibit '
        'multifamily and cannot be amended (typically requires 75% supermajority vote), the project is '
        'NOT VIABLE for the intended use.')

    doc.add_page_break()

    # =========================================================================
    # SECTION 8: DEVELOPMENT SCENARIOS
    # =========================================================================
    add_styled_heading(doc, 'Section 8: Development Scenarios', level=1)

    add_body_text(doc, 'Scenario A: 11-Unit (Net Buildable Density)', bold=True)
    sc_a = [
        ['Units', '11'],
        ['Building Configuration', 'Single 3-story building (~4,333 SF footprint)'],
        ['Building Zone', 'Eastern portion of parcel (outside wellhead easement)'],
        ['Net Buildable Area', '11,325 SF (0.260 acres) after setbacks + easement'],
        ['Average Unit Size', '900 SF'],
        ['Parking', '22 surface spaces in wellhead easement zone (west portion)'],
        ['Open Space', 'Wellhead zone remainder (~12,700 SF) as passive recreation'],
        ['Estimated Cost', '$3,786,326 ($344,211/unit)'],
    ]
    add_table(doc, ['Parameter', 'Value'], sc_a, [2.5, 4.5])

    add_body_text(doc, '')
    add_body_text(doc, 'Scenario B: 22-Unit (Maximum Gross Density)', bold=True)
    sc_b = [
        ['Units', '22'],
        ['Building Configuration', 'Two 4-story buildings (~3,063 SF footprint each)'],
        ['Building Zone', 'Full parcel (requires wellhead easement resolution)'],
        ['Net Buildable Area', '9,861 SF per building zone after 40-ft setbacks'],
        ['Average Unit Size', '850 SF'],
        ['Parking', '44 surface spaces'],
        ['Open Space', 'Dedicated preserve area + rooftop amenities'],
        ['Estimated Cost', '$7,258,503 ($329,932/unit)'],
        ['Timeline Constraint', 'Requires wellhead easement vacation (~10 years)'],
    ]
    add_table(doc, ['Parameter', 'Value'], sc_b, [2.5, 4.5])

    add_body_text(doc, '')
    add_body_text(doc, 'Setback Scenarios by Building Height', bold=True)
    setback_data = [
        ['2-story (25 ft)', '25 ft', '25 ft', '25 ft', '25 ft', '14,555 SF', '30.0%'],
        ['3-story (35 ft)', '35 ft', '35 ft', '35 ft', '35 ft', '11,325 SF', '23.3%'],
        ['4-story (40 ft)', '40 ft', '40 ft', '40 ft', '40 ft', '9,861 SF', '20.3%'],
    ]
    add_table(doc, ['Height', 'Front', 'Side', 'Rear', 'Side', 'Net Buildable', '% Parcel'],
              setback_data, [1.2, 0.8, 0.8, 0.8, 0.8, 1.2, 0.8])

    add_body_text(doc, '')
    add_body_text(doc, 'Note: RM-15 rule requires setbacks equal to building height when height exceeds '
                  'minimum setback. At 3+ stories, all setbacks equal building height, severely constraining '
                  'the buildable envelope.', italic=True)

    doc.add_page_break()

    # =========================================================================
    # SECTION 9: COST ESTIMATE
    # =========================================================================
    add_styled_heading(doc, 'Section 9: Cost Estimate', level=1)

    add_body_text(doc, '11-Unit Development Cost Summary', bold=True)
    cost_11 = [
        ['Land Acquisition', '$55,000', '1.5%'],
        ['Hard Costs (construction + site)', '$3,011,125', '79.5%'],
        ['Soft Costs (design, engineering, permits)', '$357,684', '9.4%'],
        ['Environmental (mitigation, studies)', '$77,310', '2.0%'],
        ['Impact Fees', '$130,207', '3.4%'],
        ['Financing', '$155,000', '4.1%'],
        [{'text': 'TOTAL', 'bold': True}, {'text': '$3,786,326', 'bold': True}, {'text': '100%', 'bold': True}],
        [{'text': 'Per Unit', 'bold': True}, {'text': '$344,211', 'bold': True}, {'text': '', 'bold': False}],
    ]
    add_table(doc, ['Category', 'Amount', '% of Total'], cost_11, [3.5, 1.8, 1.2])

    add_body_text(doc, '')
    add_body_text(doc, '22-Unit Development Cost Summary', bold=True)
    cost_22 = [
        ['Land Acquisition', '$55,000', '0.8%'],
        ['Hard Costs (construction + site)', '$6,023,688', '83.0%'],
        ['Soft Costs', '$527,841', '7.3%'],
        ['Environmental', '$101,560', '1.4%'],
        ['Impact Fees', '$260,414', '3.6%'],
        ['Financing', '$290,000', '4.0%'],
        [{'text': 'TOTAL', 'bold': True}, {'text': '$7,258,503', 'bold': True}, {'text': '100%', 'bold': True}],
        [{'text': 'Per Unit', 'bold': True}, {'text': '$329,932', 'bold': True}, {'text': '', 'bold': False}],
    ]
    add_table(doc, ['Category', 'Amount', '% of Total'], cost_22, [3.5, 1.8, 1.2])

    add_body_text(doc, '')
    add_body_text(doc, 'Financial Feasibility (Rental)', bold=True)
    rental_data = [
        ['Weighted Avg Rent (11 units)', '$1,575/month'],
        ['Gross Annual Rent', '$207,900'],
        ['Vacancy Rate', '7%'],
        ['Effective Gross Income', '$193,347'],
        ['Operating Expenses (40%)', '$77,339'],
        ['Net Operating Income (NOI)', '$116,008'],
        ['Yield on Cost', '3.06%'],
        ['Stabilized Value (6% cap)', '$1,933,467'],
        [{'text': 'Value Gap', 'bold': True, 'color': RED}, {'text': '-$1,852,859', 'bold': True, 'color': RED}],
    ]
    add_table(doc, ['Metric', 'Value'], rental_data, [3.5, 3.5])

    add_body_text(doc, '')
    add_body_text(doc,
        'The project does NOT pencil as a rental development. Total development cost ($3.79M) exceeds '
        'stabilized value ($1.93M) by approximately 2x. For-sale condos at $275K/unit also fail to cover '
        'development costs (net revenue $2.78M vs. cost $3.79M, a $1.0M shortfall).', italic=True)

    doc.add_page_break()

    # =========================================================================
    # SECTION 10: PERMIT ROADMAP
    # =========================================================================
    add_styled_heading(doc, 'Section 10: Permit Roadmap', level=1)

    permit_table_data = [
        ['PUD Preliminary Dev Plan', 'Palm Bay P&Z', '3-6 months', 'Public hearings required'],
        ['PUD Final Dev Plan', 'Palm Bay P&Z', '2-4 months', 'After PDP approved'],
        ['Site Plan Approval', 'Palm Bay DRC', '4-8 weeks', 'Concurrent with FDP'],
        ['Floodplain Dev Permit', 'Palm Bay Building', '2-4 weeks', 'After BFE established'],
        ['SJRWMD ERP', 'SJRWMD', '3-6 months', 'Longest lead item'],
        ['Section 404 (state)', 'FDEP', 'Concurrent', 'Joint with ERP'],
        ['FEMA LOMA/LOMR-F', 'FEMA', '60 days', 'Optional — if fill strategy'],
        ['MTWCD Stormwater', 'MTWCD', '4-8 weeks', 'Concurrent with ERP'],
        ['Building Permit', 'Palm Bay Building', '4-8 weeks', 'After all approvals'],
        ['Utility Connections', 'Palm Bay Utilities', '2-6 weeks', 'Extension may be needed'],
        ['Driveway Permit', 'Palm Bay PW', '2-4 weeks', 'Chapter 179'],
        ['HOA ARC Approval', 'Sandy Pines HOA', '1-3 months', 'CRITICAL — potential fatal'],
        ['FWC Gopher Tortoise', 'FWC', '4-8 weeks', 'If burrows found'],
    ]
    add_table(doc, ['Permit', 'Agency', 'Timeline', 'Notes'], permit_table_data, [2.2, 1.8, 1.3, 1.7])

    add_body_text(doc, '')
    add_body_text(doc, 'Estimated Total Timeline: 18-24 months (optimistic to realistic) from project '
                  'initiation to Certificate of Occupancy. Worst case: 36 months with appeals/redesign.', bold=True)

    doc.add_page_break()

    # =========================================================================
    # SECTION 11: RISK MATRIX (LANDSCAPE)
    # =========================================================================
    # Add a new section with landscape orientation
    new_section = doc.add_section(WD_ORIENT.LANDSCAPE)
    new_section.orientation = WD_ORIENT.LANDSCAPE
    new_section.page_width = Inches(11)
    new_section.page_height = Inches(8.5)
    new_section.left_margin = Cm(2.0)
    new_section.right_margin = Cm(2.0)

    add_styled_heading(doc, 'Section 11: Risk Matrix', level=1)
    add_body_text(doc, 'Constraint Severity Assessment — Color Coded by Risk Level', bold=True)

    risk_table = doc.add_table(rows=6, cols=7)
    risk_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    risk_table.style = 'Table Grid'

    # Headers
    headers = ['Constraint', 'Severity', 'Coverage', 'Est. Cost', 'Timeline', 'Mitigation', 'Fatal?']
    for i, h in enumerate(headers):
        style_table_cell(risk_table.rows[0].cells[i], h, WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    style_table_header(risk_table.rows[0])

    # Row 1: HOA CC&Rs - RED
    r1 = risk_table.rows[1]
    risk_rows_data = [
        ('HOA CC&Rs — Multifamily\nLikely Prohibited', 'POTENTIAL\nFATAL', '100%', '$5K legal\nreview',
         '1-2 months', 'CC&R amendment (75%\nvote) or legal separation', 'YES'),
        ('No Sewer/Water\nService Available', 'HIGH', '100%', '$5K-$25K\nextension',
         '2-6 months', 'Sewer main extension from\nadjacent parcel (2153)', 'NO'),
        ('FEMA Zone A Flood\n(No BFE)', 'HIGH', '100%', '$5K-$25K study\n+ 17.5% premium',
         '6-12 weeks\nstudy', 'BFE determination, elevated\nconstruction, LOMA/LOMR-F', 'NO'),
        ('NWI Wetland PEM1Cd\n(Freshwater Emergent)', 'HIGH', '100%*', '$43K-$108K\nmitigation',
         '4-6 months\nERP', 'Mitigation bank credits\n(Farmton), minimize impact', 'NO'),
        ('Wellhead Protection\nEasement (200-ft radius)', 'HIGH', '47%', '$0 (wait)\nor redesign',
         '~10 years', 'Phase 1 outside easement;\nPhase 2 after vacation', 'NO'),
    ]

    severity_colors = ['CC0000', 'FF8C00', 'FF8C00', 'FF8C00', 'FF8C00']  # RED, ORANGE x4

    for r_idx, (data, color) in enumerate(zip(risk_rows_data, severity_colors)):
        row = risk_table.rows[r_idx + 1]
        for c_idx, val in enumerate(data):
            style_table_cell(row.cells[c_idx], val, WD_ALIGN_PARAGRAPH.LEFT, size=8)
        # Color the severity cell
        set_cell_shading(row.cells[1], color)
        for p in row.cells[1].paragraphs:
            for run in p.runs:
                run.font.color.rgb = WHITE

    add_body_text(doc, '')
    add_body_text(doc, '* NWI coverage is 100% of parcel per polygon query, but actual jurisdictional '
                  'wetland boundary requires professional delineation. Partially drained (d) modifier may '
                  'reduce actual jurisdictional area.', italic=True)

    add_body_text(doc, '')
    add_body_text(doc, 'Combined Risk Assessment:', bold=True)
    add_body_text(doc,
        'VERY HIGH. Five overlapping constraints create a "constraint stacking" effect where each issue '
        'compounds the others. The HOA CC&R constraint is potentially fatal and must be resolved first. '
        'Even if all constraints can be mitigated, the combined cost ($77K-$203K environmental + $130K-$260K '
        'impact fees + $30K-$66K permit fees) far exceeds the economic capacity of this 1.1-acre parcel.')

    # Switch back to portrait
    new_section2 = doc.add_section(WD_ORIENT.PORTRAIT)
    new_section2.orientation = WD_ORIENT.PORTRAIT
    new_section2.page_width = Inches(8.5)
    new_section2.page_height = Inches(11)
    new_section2.left_margin = Cm(2.5)
    new_section2.right_margin = Cm(2.5)

    # =========================================================================
    # SECTION 12: RECOMMENDATION
    # =========================================================================
    add_styled_heading(doc, 'Section 12: Recommendation', level=1)

    p = doc.add_paragraph()
    run = p.add_run('RECOMMENDATION: HOLD / CONDITIONAL NO-GO')
    run.font.name = 'Arial'
    run.font.size = Pt(14)
    run.font.color.rgb = RED
    run.bold = True

    add_body_text(doc, '')
    add_body_text(doc, 'Primary Finding', bold=True)
    add_body_text(doc,
        'Multifamily development on this parcel is NOT ECONOMICALLY FEASIBLE under current conditions. '
        'The combination of five overlapping constraints — HOA CC&Rs, no sewer, flood zone, wetlands, '
        'and wellhead easement — creates a development cost that exceeds the stabilized property value by '
        'approximately 2x. Environmental compliance costs alone ($77K-$203K) exceed the land value ($55K).')

    add_body_text(doc, '')
    add_body_text(doc, 'Immediate Actions (Cost: $0-$50)', bold=True)
    actions = [
        ('1.', ' Obtain CC&Rs (OR Book 3636/2327) from Brevard County Clerk. This is the #1 priority. '
         'If multifamily is prohibited, no further analysis is warranted for the intended use.'),
        ('2.', ' Obtain PUD ordinance from Palm Bay City Clerk at (321) 733-3042. This determines all '
         'permitted development standards.'),
        ('3.', ' Call Palm Bay Utilities at (321) 952-3420 to confirm sewer availability and extension costs.'),
    ]
    for prefix, text in actions:
        add_bullet(doc, text, bold_prefix=prefix)

    add_body_text(doc, '')
    add_body_text(doc, 'Decision Framework', bold=True)

    decision_data = [
        ['CC&Rs prohibit multifamily AND\ncannot be amended', {'text': 'KILL', 'bold': True, 'color': RED},
         'Sell parcel or hold for alternative use'],
        ['CC&Rs allow multifamily BUT\nPUD ordinance restricts density', {'text': 'REASSESS', 'bold': True, 'color': ORANGE},
         'Evaluate SF/townhome development'],
        ['CC&Rs and PUD both allow\nmultifamily', {'text': 'PROCEED\nWITH CAUTION', 'bold': True, 'color': YELLOW_DARK},
         'Commission wetland delineation\nand flood study ($8K-$33K)'],
    ]
    add_table(doc, ['Scenario', 'Decision', 'Next Step'], decision_data, [3.0, 1.5, 2.5])

    add_body_text(doc, '')
    add_body_text(doc, 'Alternative Strategies', bold=True)
    alt_items = [
        ('Land Bank:', ' Hold parcel for 10+ years until wellhead easement is vacated, then reassess with '
         'full parcel available. Low carrying cost (vacant land taxes ~$1,000/year).'),
        ('Single-Family/Townhome:', ' Develop 2-4 townhome units consistent with HOA character. Lower '
         'environmental impact, may avoid some constraints. Requires CC&R and PUD confirmation.'),
        ('Sell:', ' List parcel at $55K-$75K. Two consecutive tax deed sales suggest limited buyer interest, '
         'but land values in Palm Bay have been rising.'),
        ('Conservation Easement:', ' Donate development rights for tax deduction. Parcel is within NWI '
         'wetland polygon — may qualify for conservation programs.'),
    ]
    for prefix, text in alt_items:
        add_bullet(doc, text, bold_prefix=prefix)

    add_body_text(doc, '')
    add_body_text(doc, '')

    # Disclaimer
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('DISCLAIMER')
    run.font.name = 'Arial'
    run.font.size = Pt(8)
    run.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(
        'This report is prepared for informational purposes only and does not constitute legal, engineering, '
        'or financial advice. All cost estimates are approximate and based on publicly available data as of '
        'February 2026. Actual costs, timelines, and regulatory requirements may differ significantly. '
        'Professional consultants (licensed engineer, surveyor, environmental scientist, real estate attorney) '
        'should be engaged before making any development decisions. The analysis relies on data from BCPAO, '
        'Palm Bay GIS, FEMA NFHL, NWI, SJRWMD, and other public sources. Data accuracy is not guaranteed. '
        'CC&R and PUD ordinance review by qualified legal counsel is essential before proceeding.')
    run.font.name = 'Arial'
    run.font.size = Pt(7)
    run.italic = True

    # Save
    output_dir = os.path.dirname(os.path.abspath(__file__))
    output_path = os.path.join(output_dir, 'Bliss_Palm_Bay_Feasibility_Report.docx')
    doc.save(output_path)
    print(f"Report saved to: {output_path}")
    return output_path

if __name__ == '__main__':
    create_report()
